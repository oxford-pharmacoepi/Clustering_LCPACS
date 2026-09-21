import csv
import gzip
from openpyxl import load_workbook

wb = load_workbook('Draft/DPA_Replication_VL_20260120.xlsx', data_only=True)
ws = wb['DPA_SNPs_FinnGen']
rows = list(ws.iter_rows(values_only=True))
header = [str(h) if h is not None else '' for h in rows[0]]
idx = {h: i for i, h in enumerate(header)}

best = {}
for r in rows[1:]:
    if not r:
        continue
    snp = r[idx['rsID']]
    logp = r[idx['LOG10P']]
    if snp is None or logp is None:
        continue
    if (snp not in best) or (logp > best[snp][idx['LOG10P']]):
        best[snp] = r

map_file = {
    'AllPCCvsGenPop': 'GWAS/pcc_vsallfil.txt.gz',
    'Subtype1vsPopCtrl': 'GWAS/clust1_vsallfil.txt.gz',
    'Subtype2vsPopCtrl': 'GWAS/clust2_vsallfil.txt.gz',
    'Subtype3vsCOVID': 'GWAS/clust3_vsno.txt.gz',
}

beta_cache = {}


def load_beta(file_path):
    if file_path in beta_cache:
        return beta_cache[file_path]
    out = {}
    try:
        with gzip.open(file_path, 'rt') as f:
            rd = csv.DictReader(f, delimiter='\t')
            for row in rd:
                s = row.get('SNP')
                b = row.get('BETA')
                if not s or not b:
                    continue
                try:
                    out[s] = float(b)
                except ValueError:
                    pass
    except FileNotFoundError:
        pass
    beta_cache[file_path] = out
    return out

records = []
for snp, r in best.items():
    analysis = r[idx['DPA_Analysis']]
    file_path = map_file.get(analysis)
    my_beta = None
    if file_path:
        my_beta = load_beta(file_path).get(snp)

    f_beta = r[idx['BETA']]
    direction_same = None
    if isinstance(f_beta, (int, float)) and isinstance(my_beta, (int, float)):
        direction_same = 1 if ((f_beta > 0 and my_beta > 0) or (f_beta < 0 and my_beta < 0)) else 0

    logp = float(r[idx['LOG10P']])
    tier = 'not_nominal'
    if logp >= 3:
        tier = 'p<=0.001'
    elif logp >= 2:
        tier = 'p<=0.01'
    elif logp >= 1.30103:
        tier = 'p<0.05'

    records.append({
        'SNP': snp,
        'analysis': analysis,
        'gene': r[idx['NearestGene']],
        'fingen_endpoint': r[idx['FinnGen_Analysis']],
        'fingen_beta': f_beta,
        'fingen_se': r[idx['SE']],
        'fingen_log10p': logp,
        'fingen_p': 10 ** (-logp),
        'my_beta': my_beta,
        'direction_same': direction_same,
        'replication_tier': tier,
    })

records.sort(key=lambda x: x['fingen_log10p'], reverse=True)

with open('Draft/fingen_replication_summary.csv', 'w', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(records[0].keys()))
    writer.writeheader()
    writer.writerows(records)

print('n_index_snps', len(records))
print('nominal_replications', sum(1 for r in records if r['fingen_log10p'] >= 1.30103))
print('p_le_0.01', sum(1 for r in records if r['fingen_log10p'] >= 2))
print('p_le_0.001', sum(1 for r in records if r['fingen_log10p'] >= 3))
print('direction_concordant_n', sum(1 for r in records if r['direction_same'] == 1))
print('direction_discordant_n', sum(1 for r in records if r['direction_same'] == 0))
print('direction_missing_n', sum(1 for r in records if r['direction_same'] is None))
